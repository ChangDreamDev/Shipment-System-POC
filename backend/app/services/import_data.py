import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.schemas import Load


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def _find_sheet(wb, names: list[str]):
    for sheet in wb.sheetnames:
        if sheet.strip().lower() in {n.lower() for n in names}:
            return wb[sheet]
    for sheet in wb.sheetnames:
        for name in names:
            if name.lower() in sheet.strip().lower():
                return wb[sheet]
    return None


def _sheet_to_text(sheet) -> str:
    lines: list[str] = []
    for row in sheet.iter_rows(values_only=True):
        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if cells:
            lines.append(" ".join(cells))
    return "\n".join(lines)


def _parse_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace(",", "").replace("$", "").strip())
    except ValueError:
        return None


def _parse_load_row(headers: dict[str, int], row: tuple) -> Load:
    def cell(name: str) -> Any:
        idx = headers.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    load_id = str(cell("load_id") or cell("id") or cell("load id") or "").strip()
    if not load_id:
        load_id = str(cell("load#") or "").strip()

    origin_city = cell("origin") or cell("origin_city") or cell("origin city")
    dest_city = cell("destination") or cell("destination_city") or cell("destination city")

    load = Load(
        load_id=load_id or f"row-{hash(row) % 100000}",
        origin_city=str(origin_city).strip() if origin_city else None,
        origin_lat=_parse_float(cell("origin_lat") or cell("origin lat")),
        origin_lon=_parse_float(cell("origin_lon") or cell("origin lng") or cell("origin lon")),
        destination_city=str(dest_city).strip() if dest_city else None,
        destination_lat=_parse_float(
            cell("destination_lat") or cell("dest_lat") or cell("destination lat")
        ),
        destination_lon=_parse_float(
            cell("destination_lon")
            or cell("dest_lon")
            or cell("destination lng")
            or cell("destination lon")
        ),
        price=_parse_float(cell("price") or cell("rate") or cell("pay")),
        weight_lbs=_parse_float(cell("weight") or cell("weight_lbs") or cell("weight lbs")),
        trailer_type=str(cell("trailer_type") or cell("trailer") or cell("equipment") or "")
        .strip()
        or None,
    )

    complete, reason = _mark_incomplete(load)
    load.incomplete = not complete
    load.incomplete_reason = reason
    return load


def _mark_incomplete(load: Load) -> tuple[bool, str | None]:
    missing = []
    if load.price is None:
        missing.append("price")
    if load.destination_lat is None or load.destination_lon is None:
        missing.append("destination")
    if missing:
        return False, f"Missing {', '.join(missing)}"
    return True, None


def _build_header_map(header_row: tuple) -> dict[str, int]:
    headers: dict[str, int] = {}
    for i, cell in enumerate(header_row):
        key = _norm_header(cell)
        if key:
            headers[key] = i
    return headers


def _alias_headers(headers: dict[str, int]) -> dict[str, int]:
    """Map flexible column names to canonical keys."""
    aliases = {
        "load_id": ["load id", "load_id", "id", "load#", "load number", "load"],
        "origin_lat": ["origin_lat", "origin lat", "origin latitude", "orig lat"],
        "origin_lon": ["origin_lon", "origin lng", "origin lon", "origin longitude", "orig lon"],
        "destination_lat": [
            "destination_lat",
            "dest_lat",
            "destination lat",
            "dest lat",
            "destination latitude",
        ],
        "destination_lon": [
            "destination_lon",
            "dest_lon",
            "destination lng",
            "destination lon",
            "dest lon",
            "destination longitude",
        ],
        "price": ["price", "rate", "pay", "total pay", "linehaul"],
        "weight": ["weight", "weight_lbs", "weight lbs", "weight (lbs)"],
        "trailer_type": ["trailer_type", "trailer", "equipment", "trailer type"],
        "origin": ["origin", "origin_city", "origin city", "pickup"],
        "destination": ["destination", "destination_city", "destination city", "delivery"],
    }
    result: dict[str, int] = {}
    for canonical, names in aliases.items():
        for name in names:
            if name in headers:
                result[canonical] = headers[name]
                break
    return result


def import_from_xlsx(path: Path) -> tuple[str, list[Load]]:
    wb = load_workbook(path, read_only=True, data_only=True)

    conv_sheet = _find_sheet(wb, ["Sample Conversation", "Conversation", "Transcript"])
    loads_sheet = _find_sheet(wb, ["Loads", "Load Board", "Available Loads"])

    if conv_sheet is None:
        raise ValueError("Could not find 'Sample Conversation' sheet in workbook")
    if loads_sheet is None:
        raise ValueError("Could not find 'Loads' sheet in workbook")

    transcript = _sheet_to_text(conv_sheet)

    rows = list(loads_sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Loads sheet is empty")

    header_map = _alias_headers(_build_header_map(rows[0]))
    loads: list[Load] = []
    for row in rows[1:]:
        if not any(row):
            continue
        load = _parse_load_row(header_map, row)
        if load.origin_city or load.price or load.trailer_type:
            loads.append(load)

    return transcript, loads
